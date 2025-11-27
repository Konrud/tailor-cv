"""File processing service for extracting text from PDF, DOCX, and TXT files.

This module provides text extraction functionality for supported file formats
using PyMuPDF (PDF), python-docx (DOCX), and built-in file reading (TXT).
"""

import io
from typing import BinaryIO

import docx
import pymupdf as fitz  # PyMuPDF

from src.utils.errors import FileProcessingError


def extract_text_from_pdf(file_content: bytes | BinaryIO) -> tuple[str, int]:
    """Extract text content from a PDF file using PyMuPDF.

    Args:
        file_content: PDF file content as bytes or file-like object

    Returns:
        tuple: (extracted_text, page_count)

    Raises:
        FileProcessingError: If PDF processing fails
    """
    try:
        # Convert bytes to file-like object if necessary
        if isinstance(file_content, bytes):
            file_stream = io.BytesIO(file_content)
        else:
            file_stream = file_content

        # Open PDF document
        pdf_document = fitz.open(stream=file_stream, filetype="pdf")
        page_count = len(pdf_document)

        # Extract text from all pages
        text_content = []
        for page_num in range(page_count):
            page = pdf_document[page_num]
            text = page.get_text()
            text_content.append(text)

        pdf_document.close()

        # Combine all pages
        full_text = "\n\n".join(text_content).strip()

        if not full_text:
            raise FileProcessingError(
                "No text could be extracted from the PDF. Please ensure the PDF contains selectable text (not just images).",
                details="PDF appears to be empty or contains only images",
            )

        return full_text, page_count

    except fitz.FileDataError as e:
        raise FileProcessingError(
            "Invalid or corrupted PDF file",
            details=str(e),
        ) from e
    except Exception as e:
        raise FileProcessingError(
            "Failed to extract text from PDF",
            details=str(e),
        ) from e


def extract_text_from_docx(file_content: bytes | BinaryIO) -> tuple[str, int]:
    """Extract text content from a DOCX file using python-docx.

    Args:
        file_content: DOCX file content as bytes or file-like object

    Returns:
        tuple: (extracted_text, estimated_page_count)

    Raises:
        FileProcessingError: If DOCX processing fails
    """
    try:
        # Convert bytes to file-like object if necessary
        if isinstance(file_content, bytes):
            file_stream = io.BytesIO(file_content)
        else:
            file_stream = file_content

        # Open DOCX document
        doc = docx.Document(file_stream)

        # Extract text from all paragraphs
        text_content = []
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_content.append(paragraph.text)

        # Extract text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text_content.append(cell.text)

        full_text = "\n".join(text_content).strip()

        if not full_text:
            raise FileProcessingError(
                "No text could be extracted from the DOCX file",
                details="DOCX appears to be empty",
            )

        # Estimate page count (roughly 500 words per page)
        word_count = len(full_text.split())
        estimated_pages = max(1, (word_count + 499) // 500)

        return full_text, estimated_pages

    except docx.opc.exceptions.PackageNotFoundError as e:
        raise FileProcessingError(
            "Invalid or corrupted DOCX file",
            details=str(e),
        ) from e
    except Exception as e:
        raise FileProcessingError(
            "Failed to extract text from DOCX",
            details=str(e),
        ) from e


def extract_text_from_txt(file_content: bytes | BinaryIO) -> tuple[str, int]:
    """Extract text content from a TXT file.

    Args:
        file_content: TXT file content as bytes or file-like object

    Returns:
        tuple: (extracted_text, estimated_page_count)

    Raises:
        FileProcessingError: If TXT processing fails
    """
    try:
        # Read content
        if isinstance(file_content, bytes):
            text = file_content.decode("utf-8", errors="ignore")
        else:
            text = file_content.read()
            if isinstance(text, bytes):
                text = text.decode("utf-8", errors="ignore")

        full_text = text.strip()

        if not full_text:
            raise FileProcessingError(
                "Text file is empty",
                details="No content found in TXT file",
            )

        # Estimate page count (roughly 500 words per page)
        word_count = len(full_text.split())
        estimated_pages = max(1, (word_count + 499) // 500)

        return full_text, estimated_pages

    except UnicodeDecodeError as e:
        raise FileProcessingError(
            "Failed to decode text file. Please ensure the file is in UTF-8 encoding.",
            details=str(e),
        ) from e
    except Exception as e:
        raise FileProcessingError(
            "Failed to extract text from TXT file",
            details=str(e),
        ) from e


def extract_text(file_content: bytes | BinaryIO, file_type: str) -> tuple[str, int]:
    """Extract text from a file based on its type.

    This is the main entry point for text extraction.

    Args:
        file_content: File content as bytes or file-like object
        file_type: File type ("pdf", "docx", or "txt")

    Returns:
        tuple: (extracted_text, page_count)

    Raises:
        FileProcessingError: If extraction fails
        ValueError: If file type is not supported
    """
    extractors = {
        "pdf": extract_text_from_pdf,
        "docx": extract_text_from_docx,
        "txt": extract_text_from_txt,
    }

    extractor = extractors.get(file_type.lower())
    if not extractor:
        raise ValueError(f"Unsupported file type: {file_type}")

    return extractor(file_content)

